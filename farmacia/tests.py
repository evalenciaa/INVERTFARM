import json
import re
import threading
from datetime import date, timedelta
from io import BytesIO

from django.contrib.auth.models import Group, Permission
from django.core import mail
from django.core.exceptions import ValidationError
from django.db import IntegrityError, close_old_connections, transaction
from django.test import (
    Client, TestCase, TransactionTestCase, override_settings, skipUnlessDBFeature,
)
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from .models import (
    Almacen, DetalleEntrada, DetalleSalidaTransferencia, Entrada,
    CPMMedicamento, FolioConsecutivo, FuenteFinanciamiento, Institucion, Lote,
    Medicamento, Paciente, Presentacion, Receta, RecetaMedicamento,
    MedicamentoNoSurtido, MedicamentoNoDisponibleTransferencia,
    SalidaTransferencia, UsuarioPersonalizado,
)
from .cpm import actualizar_cpm_medicamento, calcular_estado_inventario, periodo_cpm
from .services import descontar_lotes, surtir_fefo
from .tasks import verificar_alertas_cpm


def conceder(user, *codenames):
    user.user_permissions.add(*Permission.objects.filter(codename__in=codenames))
    for cache_name in ('_perm_cache', '_user_perm_cache', '_group_perm_cache'):
        user.__dict__.pop(cache_name, None)


class DatosFarmaciaMixin:
    def crear_catalogos(self):
        self.presentacion = Presentacion.objects.create(nombre='UNIDAD', unidades_por_caja=1)
        self.medicamento = Medicamento.objects.create(
            clave='010.000.0001', descripcion='Medicamento de prueba',
            presentacion=self.presentacion, costo='12.50'
        )
        self.almacen = Almacen.objects.create(
            codigo='ALM-01', nombre='Almacén prueba', direccion='Prueba'
        )
        self.fuente = FuenteFinanciamiento.objects.create(codigo='FF-01', nombre='Fuente prueba')

    def crear_lote(self, identificador='LOT-PRUEBA', existencia=10, dias=365):
        return Lote.objects.create(
            id=identificador, medicamento=self.medicamento,
            lote_codigo=identificador, fecha_caducidad=date.today() + timedelta(days=dias),
            existencia=existencia, presentacion=self.presentacion, costo_unitario='12.50'
        )


class LoginFuncionalTests(TestCase):
    def setUp(self):
        self.usuario = UsuarioPersonalizado.objects.create_user(
            username='usuario-login',
            password='PruebaSegura123!',
            rol='FARMACIA',
        )

    def test_login_muestra_formulario_funcional_y_estado_seguro(self):
        response = self.client.get(reverse('login'), {'next': reverse('farmacia')})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="login-shell"')
        self.assertContains(response, 'name="username"')
        self.assertContains(response, 'name="password"')
        self.assertContains(response, 'name="csrfmiddlewaretoken"')
        self.assertContains(response, f'value="{reverse("farmacia")}"')
        self.assertContains(response, 'Acceso protegido')
        self.assertNotContains(response, 'Intento de acceso incorrecto')

    def test_login_incorrecto_conserva_usuario_y_muestra_error(self):
        response = self.client.post(reverse('login'), {
            'username': self.usuario.username,
            'password': 'incorrecta',
        })

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'value="{self.usuario.username}"')
        self.assertContains(response, 'Usuario o contraseña incorrectos')

    def test_login_correcto_conserva_redireccion_segura(self):
        destino = reverse('farmacia')
        response = self.client.post(
            f'{reverse("login")}?next={destino}',
            {
                'username': self.usuario.username,
                'password': 'PruebaSegura123!',
                'next': destino,
            },
        )

        self.assertRedirects(response, destino, fetch_redirect_response=False)
        self.assertEqual(
            str(self.client.session.get('_auth_user_id')),
            str(self.usuario.pk),
        )


class CPMCalculadoTests(DatosFarmaciaMixin, TestCase):
    def setUp(self):
        self.crear_catalogos()
        self.paciente = Paciente.objects.create(
            nombre_completo='Paciente CPM',
            curp='CPMPRUEBA123456789',
            fecha_nacimiento=date(1990, 1, 1),
        )

    def registrar_consumo(self, folio, fecha_surtido, cantidad):
        receta = Receta.objects.create(
            id_folio=folio,
            paciente=self.paciente,
            fecha_emision=fecha_surtido,
            fecha_surtido=fecha_surtido,
            estado='completa',
            origen='hospitalizacion_adultos',
        )
        return RecetaMedicamento.objects.create(
            receta=receta,
            medicamento=self.medicamento,
            cantidad_solicitada=cantidad,
            cantidad_surtida=cantidad,
        )

    def test_cpm_suma_recetas_y_colectivos_en_ventana_movil_de_tres_meses(self):
        referencia = date(2026, 9, 18)
        inicio, fin = periodo_cpm(referencia)
        self.assertEqual(inicio, date(2026, 6, 18))
        self.assertEqual(fin, referencia)

        self.registrar_consumo('REC-CPM-1', date(2026, 7, 1), 4)
        # Los colectivos completados también generan RecetaMedicamento.
        self.registrar_consumo('COL-CPM-1', date(2026, 8, 15), 3)
        self.registrar_consumo('REC-CPM-2', referencia, 3)
        self.registrar_consumo('REC-ANTIGUA', inicio, 99)

        cpm = actualizar_cpm_medicamento(self.medicamento.id, referencia)

        self.assertEqual(cpm.valor, 4)  # ceil((4 + 3 + 3) / 3)
        self.assertIsNone(cpm.actualizado_por)

    @override_settings(ALERTAS_STOCK_DESTINATARIOS=['alertas@example.com'])
    def test_alerta_recalcula_cpm_y_compara_con_existencia_total(self):
        hoy = date.today()
        self.registrar_consumo('REC-ALERTA-CPM', hoy, 18)
        self.crear_lote(existencia=3)
        CPMMedicamento.objects.update_or_create(
            medicamento=self.medicamento,
            defaults={'valor': 1},
        )

        resultado = verificar_alertas_cpm()

        self.assertEqual(
            CPMMedicamento.objects.get(medicamento=self.medicamento).valor,
            6,
        )
        self.assertEqual(resultado, 'Alertas enviadas: 1')
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn('Existencia: 3 | CPM: 6 | 50.0%', mail.outbox[0].body)

    def test_estado_usa_cpm_stock_maximo_y_excedente(self):
        casos = (
            (0, 'desabasto', 0, 0.0),
            (19, 'bajo', 0, 38.0),
            (20, 'adecuado', 0, 40.0),
            (50, 'adecuado', 0, 100.0),
            (51, 'excedente', 1, 102.0),
        )
        for existencia, estado, excedente, porcentaje in casos:
            with self.subTest(existencia=existencia):
                resultado = calcular_estado_inventario(existencia, 20)
                self.assertEqual(resultado['stock_maximo'], 50)
                self.assertEqual(resultado['estado'], estado)
                self.assertEqual(resultado['excedente'], excedente)
                self.assertEqual(resultado['porcentaje'], porcentaje)

    def test_inventario_general_incluye_lotes_en_desabasto(self):
        self.crear_lote(existencia=0)
        CPMMedicamento.objects.update_or_create(
            medicamento=self.medicamento,
            defaults={'valor': 20},
        )
        admin = UsuarioPersonalizado.objects.create_superuser(
            username='admin-cpm', password='PruebaSegura123!', rol='ADMIN'
        )
        self.client.force_login(admin)

        response = self.client.get(reverse('inv_gene_f'))

        self.assertEqual(response.status_code, 200)
        item = response.context['inventario'][0]
        self.assertEqual(item['estado'], 'desabasto')
        self.assertEqual(item['stock_maximo'], 50)
        self.assertContains(response, 'Desabasto')
        self.assertContains(response, 'PDF Excedentes')
        self.assertContains(
            response, reverse('exportar_inventario_general_excedentes_pdf')
        )

    def test_reporte_de_excedentes_filtra_con_la_formula_de_stock_maximo(self):
        from .views.reporte_views import _obtener_datos_inventario_general

        self.crear_lote(existencia=51)
        CPMMedicamento.objects.update_or_create(
            medicamento=self.medicamento,
            defaults={'valor': 20},
        )
        otro = Medicamento.objects.create(
            clave='010.000.0002', descripcion='Medicamento adecuado',
            presentacion=self.presentacion, costo='10.00',
        )
        Lote.objects.create(
            id='LOT-ADECUADO', medicamento=otro, lote_codigo='LOT-ADECUADO',
            fecha_caducidad=date.today() + timedelta(days=365),
            existencia=50, presentacion=self.presentacion, costo_unitario='10.00',
        )
        CPMMedicamento.objects.update_or_create(
            medicamento=otro,
            defaults={'valor': 20},
        )

        datos = _obtener_datos_inventario_general(solo_excedentes=True)

        self.assertEqual(len(datos), 1)
        self.assertEqual(datos[0]['medicamento__id'], self.medicamento.id)
        self.assertEqual(datos[0]['stock_maximo'], 50)
        self.assertEqual(datos[0]['excedente'], 1)

    def test_pdf_general_divide_inventario_extenso_en_varias_paginas(self):
        admin = UsuarioPersonalizado.objects.create_superuser(
            username='admin-pdf-general', password='PruebaSegura123!', rol='ADMIN'
        )
        for numero in range(45):
            medicamento = Medicamento.objects.create(
                clave=f'010.999.{numero:04d}',
                descripcion=(
                    f'Medicamento de prueba {numero} con descripción extensa para '
                    'comprobar que la tabla se divide correctamente entre páginas.'
                ),
                presentacion=self.presentacion,
                costo='10.00',
            )
            Lote.objects.create(
                id=f'LOT-PDF-{numero:04d}', medicamento=medicamento,
                lote_codigo=f'LOT-PDF-{numero:04d}',
                fecha_caducidad=date.today() + timedelta(days=365),
                existencia=numero + 1, presentacion=self.presentacion,
                costo_unitario='10.00',
            )
        self.client.force_login(admin)

        response = self.client.get(reverse('exportar_inventario_general_pdf'))

        self.assertEqual(response.status_code, 200, response.content[:500])
        self.assertEqual(response['Content-Type'], 'application/pdf')
        paginas = len(re.findall(rb'/Type\s*/Page\b', response.content))
        self.assertGreaterEqual(paginas, 3)


class BalanceAntibioticosAwareTests(DatosFarmaciaMixin, TestCase):
    def setUp(self):
        self.crear_catalogos()
        self.admin = UsuarioPersonalizado.objects.create_superuser(
            username='admin-aware', password='PruebaSegura123!', rol='ADMIN'
        )
        self.paciente = Paciente.objects.create(
            nombre_completo='Paciente AWaRe',
            curp='AWAREPRUEBA1234567',
            fecha_nacimiento=date(1990, 1, 1),
        )
        self.client.force_login(self.admin)

    def crear_antibiotico(self, clave, categoria, gramos, valor_atc):
        return Medicamento.objects.create(
            clave=clave,
            descripcion=f'Antibiótico {categoria}',
            presentacion=self.presentacion,
            costo='10.00',
            es_antibiotico=True,
            codigo_atc=f'J01{clave[-2:]}',
            categoria_aware=categoria,
            gramos_por_pieza=gramos,
            valor_atc=valor_atc,
        )

    def registrar_consumo(self, folio, medicamento, cantidad):
        receta = Receta.objects.create(
            id_folio=folio,
            paciente=self.paciente,
            fecha_emision=date(2026, 7, 10),
            fecha_surtido=date(2026, 7, 10),
            estado='completa',
            origen='hospitalizacion_adultos',
            surtido_por=self.admin,
        )
        RecetaMedicamento.objects.create(
            receta=receta,
            medicamento=medicamento,
            cantidad_solicitada=cantidad,
            cantidad_surtida=cantidad,
        )

    def test_balance_suma_ddd_y_calcula_porcentaje_por_grupo(self):
        access = self.crear_antibiotico(
            '010.000.1001', 'Access', '2.0000', '0.5000'
        )
        watch = self.crear_antibiotico(
            '010.000.1002', 'Watch', '1.0000', '1.0000'
        )
        self.crear_antibiotico(
            '010.000.1003', 'Reserve', '1.0000', '1.0000'
        )
        self.registrar_consumo('REC-AWARE-1', access, 10)  # 40 DDD
        self.registrar_consumo('COL-AWARE-1', watch, 10)   # 10 DDD

        response = self.client.get(
            reverse('inventario_antibioticos'), {'mes': '2026-07'}
        )

        self.assertEqual(response.status_code, 200)
        resumen = {
            fila['categoria']: fila
            for fila in response.context['resumen_aware']
        }
        self.assertEqual(resumen['Access']['sumatoria'], 40)
        self.assertEqual(resumen['Access']['porcentaje'], 80)
        self.assertEqual(resumen['Watch']['sumatoria'], 10)
        self.assertEqual(resumen['Watch']['porcentaje'], 20)
        self.assertEqual(resumen['Reserve']['sumatoria'], 0)
        self.assertEqual(response.context['total_ddd_antibioticos'], 50)
        self.assertContains(response, 'tipo-grafica-aware')
        self.assertContains(response, 'Balance mensual de consumo AWaRe')

    def test_balance_sin_consumo_devuelve_porcentajes_en_cero(self):
        self.crear_antibiotico(
            '010.000.1004', 'Access', '1.0000', '1.0000'
        )

        response = self.client.get(
            reverse('inventario_antibioticos'), {'mes': '2026-07'}
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['total_ddd_antibioticos'], 0)
        self.assertTrue(all(
            fila['porcentaje'] == 0
            for fila in response.context['resumen_aware']
        ))

    def test_reportes_generales_incluyen_balance_y_grafica(self):
        access = self.crear_antibiotico(
            '010.000.1101', 'Access', '2.0000', '0.5000'
        )
        watch = self.crear_antibiotico(
            '010.000.1102', 'Watch', '1.0000', '1.0000'
        )
        self.registrar_consumo('REC-AWARE-EXCEL-1', access, 10)
        self.registrar_consumo('REC-AWARE-EXCEL-2', watch, 10)
        parametros = {'mes': '2026-07'}

        excel = self.client.get(
            reverse('exportar_antibioticos_excel'), parametros
        )
        self.assertEqual(excel.status_code, 200, excel.content[:500])
        workbook = load_workbook(BytesIO(excel.content), read_only=False)
        worksheet = workbook['Antibioticos']
        contenido = '\n'.join(
            str(valor)
            for fila in worksheet.iter_rows(values_only=True)
            for valor in fila
            if valor is not None
        )
        self.assertIn('BALANCE DE CONSUMO POR CLASIFICACIÓN AWaRe', contenido)
        self.assertEqual(len(worksheet._charts), 1)

        pdf = self.client.get(
            reverse('exportar_antibioticos_pdf'), parametros
        )
        self.assertEqual(pdf.status_code, 200, pdf.content[:500])
        paginas = len(re.findall(rb'/Type\s*/Page\b', pdf.content))
        self.assertGreaterEqual(paginas, 2)
        self.assertIn('Antibioticos_202607.pdf', pdf['Content-Disposition'])

    def test_pdfs_individuales_filtran_aware_y_no_incluyen_grafica(self):
        medicamentos = {
            'access': self.crear_antibiotico(
                '010.000.1201', 'Access', '1.0000', '1.0000'
            ),
            'watch': self.crear_antibiotico(
                '010.000.1202', 'Watch', '1.0000', '1.0000'
            ),
            'reserve': self.crear_antibiotico(
                '010.000.1203', 'Reserve', '1.0000', '1.0000'
            ),
        }
        for indice, medicamento in enumerate(medicamentos.values(), start=1):
            self.registrar_consumo(
                f'REC-AWARE-PDF-{indice}', medicamento, indice
            )

        nombres_url = {
            'access': 'exportar_antibioticos_access_pdf',
            'watch': 'exportar_antibioticos_watch_pdf',
            'reserve': 'exportar_antibioticos_reserve_pdf',
        }
        for categoria, nombre_url in nombres_url.items():
            with self.subTest(categoria=categoria):
                response = self.client.get(
                    reverse(nombre_url), {'mes': '2026-07'}
                )
                self.assertEqual(response.status_code, 200, response.content[:500])
                self.assertIn(
                    f'Antibioticos_{categoria.title()}_202607.pdf',
                    response['Content-Disposition'],
                )
                paginas = len(re.findall(rb'/Type\s*/Page\b', response.content))
                self.assertEqual(paginas, 1)


class PaginacionPDFTests(DatosFarmaciaMixin, TestCase):
    def setUp(self):
        self.crear_catalogos()
        self.admin = UsuarioPersonalizado.objects.create_superuser(
            username='admin-pdf-auditoria',
            password='PruebaSegura123!',
            rol='ADMIN',
        )
        self.paciente = Paciente.objects.create(
            nombre_completo='Paciente con nombre extenso para auditoría de PDF',
            curp='PDFAUDITORIA123456',
            fecha_nacimiento=date(1990, 1, 1),
        )
        self.lote = self.crear_lote(existencia=500)
        self.institucion = Institucion.objects.create(
            codigo='HOSP-PDF',
            nombre='Hospital regional con nombre extenso para auditoría',
            tipo='HOSPITAL',
            direccion='Dirección de prueba',
        )
        self.client.force_login(self.admin)

    def test_comprobante_receta_pagina_tablas_extensas(self):
        receta = Receta.objects.create(
            id_folio='REC-PDF-EXTENSA',
            paciente=self.paciente,
            fecha_emision=date.today(),
            fecha_surtido=date.today(),
            estado='parcial',
            origen='hospitalizacion_adultos',
            surtido_por=self.admin,
        )
        RecetaMedicamento.objects.bulk_create([
            RecetaMedicamento(
                receta=receta,
                medicamento=self.medicamento,
                lote=self.lote,
                cantidad_solicitada=indice,
                cantidad_surtida=indice,
            )
            for indice in range(1, 46)
        ])
        MedicamentoNoSurtido.objects.bulk_create([
            MedicamentoNoSurtido(
                receta=receta,
                medicamento_descripcion=(
                    f'Medicamento faltante {indice} con descripción extensa '
                    'para validar el ajuste dentro de la celda'
                ),
                cantidad_solicitada=indice,
                motivo='Sin stock disponible para cubrir la cantidad solicitada',
            )
            for indice in range(1, 11)
        ])

        response = self.client.get(
            reverse('descargar_comprobante', args=[receta.id])
        )

        self.assertEqual(response.status_code, 200)
        paginas = len(re.findall(rb'/Type\s*/Page\b', response.content))
        self.assertGreaterEqual(paginas, 2)

    def test_comprobante_transferencia_pagina_tablas_extensas(self):
        transferencia = SalidaTransferencia.objects.create(
            folio='TRA-PDF-EXTENSA',
            institucion_destino=self.institucion,
            autorizado_por=self.admin,
            observaciones=(
                'Observaciones extensas para verificar que el texto se ajuste '
                'correctamente y no quede fuera del área imprimible.'
            ),
        )
        DetalleSalidaTransferencia.objects.bulk_create([
            DetalleSalidaTransferencia(
                transferencia=transferencia,
                lote=self.lote,
                cantidad=indice,
                costo_unitario='12.50',
            )
            for indice in range(1, 46)
        ])
        MedicamentoNoDisponibleTransferencia.objects.bulk_create([
            MedicamentoNoDisponibleTransferencia(
                transferencia=transferencia,
                medicamento_descripcion=(
                    f'Medicamento no disponible {indice} con descripción extensa'
                ),
                cantidad_solicitada=indice,
                motivo='Inventario insuficiente para completar la transferencia',
            )
            for indice in range(1, 11)
        ])

        response = self.client.get(
            reverse('descargar_comprobante_transferencia', args=[transferencia.id])
        )

        self.assertEqual(response.status_code, 200)
        paginas = len(re.findall(rb'/Type\s*/Page\b', response.content))
        self.assertGreaterEqual(paginas, 2)

    def test_reporte_entrada_extenso_usa_horizontal_y_pagina(self):
        payload = {
            'folio': 'ENT-PDF-EXTENSA',
            'fecha': date.today().strftime('%d/%m/%Y'),
            'tipo_entrada': 'Entrada por almacén',
            'almacen_nombre': (
                'Almacén con nombre considerablemente largo para validar '
                'el ajuste del encabezado'
            ),
            'fuente_financiamiento_nombre': (
                'Fuente de financiamiento institucional con descripción extensa'
            ),
            'proceso': 'Proceso de auditoría de PDF',
            'total': 1000,
            'items': [
                {
                    'nombre': (
                        f'Medicamento {indice} con descripción extensa para '
                        'validar paginación y ajuste dentro de la celda'
                    ),
                    'lote': f'LOTE-{indice:03d}',
                    'presentacion': 'Caja con presentación extensa',
                    'cantidad': indice,
                    'precio_unitario': 10,
                    'total': indice * 10,
                }
                for indice in range(1, 46)
            ],
        }

        response = self.client.post(
            reverse('generar_reporte_pdf'),
            data=json.dumps(payload),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200, response.content[:500])
        paginas = len(re.findall(rb'/Type\s*/Page\b', response.content))
        self.assertGreaterEqual(paginas, 2)
        self.assertIn(b'/MediaBox [ 0 0 792 612 ]', response.content)


class SeguridadEndpointsTests(DatosFarmaciaMixin, TestCase):
    def setUp(self):
        self.crear_catalogos()
        self.lote = self.crear_lote()
        self.usuario = UsuarioPersonalizado.objects.create_user(
            username='usuario', password='PruebaSegura123!', rol='FARMACIA'
        )
        self.admin = UsuarioPersonalizado.objects.create_superuser(
            username='admin', password='PruebaSegura123!', rol='ADMIN'
        )
        self.paciente = Paciente.objects.create(
            nombre_completo='Paciente Prueba', curp='CURPPRUEBA12345678',
            fecha_nacimiento=date(1990, 1, 1)
        )

    def test_datos_de_paciente_y_lote_requieren_autenticacion(self):
        urls = [
            reverse('get_paciente_info_json', args=[self.paciente.curp]),
            reverse('get_paciente_by_name', args=[self.paciente.nombre_completo]),
            reverse('buscar_lote_json', args=[self.lote.id]),
        ]
        for url in urls:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 302)

    def test_usuario_sin_permiso_no_consulta_datos_sensibles(self):
        self.client.force_login(self.usuario)
        response = self.client.get(reverse('get_paciente_info_json', args=[self.paciente.curp]))
        self.assertEqual(response.status_code, 403)

    def test_usuario_con_permiso_puede_consultar_paciente(self):
        conceder(self.usuario, 'create_salida')
        self.client.force_login(self.usuario)
        response = self.client.get(reverse('get_paciente_info_json', args=[self.paciente.curp]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['nombre_completo'], self.paciente.nombre_completo)

    def test_variante_historica_del_grupo_farmaceutico_sigue_autorizada(self):
        self.usuario.groups.add(Group.objects.create(name='Farmaceutico'))
        conceder(self.usuario, 'view_reportes')
        self.client.force_login(self.usuario)
        self.assertEqual(self.client.get(reverse('reportes_farmacia')).status_code, 200)

    def test_eliminar_medicamento_exige_permiso(self):
        self.client.force_login(self.usuario)
        response = self.client.post(
            reverse('eliminar_medicamento'),
            data=json.dumps({'medicamento_id': self.medicamento.id}),
            content_type='application/json'
        )
        self.assertEqual(response.status_code, 403)

    def test_registro_api_es_solo_para_administradores(self):
        api = APIClient()
        datos = {'username': 'externo', 'password': 'PruebaSegura123!', 'rol': 'FARMACIA'}
        self.assertEqual(api.post(reverse('api_register'), datos).status_code, 401)
        api.force_authenticate(self.admin)
        datos['username'] = 'creado_por_admin'
        response = api.post(reverse('api_register'), datos)
        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue(UsuarioPersonalizado.objects.filter(username='creado_por_admin').exists())

    def test_reportes_de_entrada_conservan_proteccion_csrf(self):
        conceder(self.usuario, 'add_entrada')
        cliente_csrf = Client(enforce_csrf_checks=True)
        cliente_csrf.force_login(self.usuario)
        response = cliente_csrf.post(
            reverse('generar_reporte_pdf'),
            data=json.dumps({'folio': 'ENT-TEST', 'items': []}),
            content_type='application/json'
        )
        self.assertEqual(response.status_code, 403)

    def test_reportes_pdf_y_excel_requieren_permiso_y_se_generan(self):
        payload = {
            'folio': 'ENT-REPORTE-1', 'fecha': date.today().isoformat(),
            'tipo_entrada': 'ALMACEN', 'items': [], 'total': 0,
        }
        self.client.force_login(self.usuario)
        self.assertEqual(
            self.client.post(
                reverse('generar_reporte_pdf'), json.dumps(payload),
                content_type='application/json'
            ).status_code,
            403,
        )
        conceder(self.usuario, 'add_entrada')
        for nombre_url, content_type in (
            ('generar_reporte_pdf', 'application/pdf'),
            ('generar_reporte_excel', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
        ):
            with self.subTest(nombre_url=nombre_url):
                response = self.client.post(
                    reverse(nombre_url), json.dumps(payload),
                    content_type='application/json'
                )
                self.assertEqual(response.status_code, 200, response.content)
                self.assertEqual(response['Content-Type'], content_type)

    def test_backups_solo_administrador_y_sin_path_traversal(self):
        self.client.force_login(self.usuario)
        self.assertNotEqual(self.client.get(reverse('panel_backups')).status_code, 200)

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse('descargar_backup', args=['..\\requirements.txt'])
        )
        self.assertEqual(response.status_code, 404)


class FlujosInventarioTests(DatosFarmaciaMixin, TestCase):
    def setUp(self):
        self.crear_catalogos()
        self.usuario = UsuarioPersonalizado.objects.create_user(
            username='farmaceutico', password='PruebaSegura123!', rol='FARMACIA'
        )

    def test_servicio_descuenta_y_rechaza_stock_insuficiente(self):
        lote = self.crear_lote(existencia=10)
        with transaction.atomic():
            descontar_lotes({lote.id: 4})
        lote.refresh_from_db()
        self.assertEqual(lote.existencia, 6)
        with self.assertRaises(ValidationError):
            with transaction.atomic():
                descontar_lotes({lote.id: 7})
        lote.refresh_from_db()
        self.assertEqual(lote.existencia, 6)

    def test_surtido_fefo_usa_primero_el_lote_mas_proximo(self):
        primero = self.crear_lote('LOT-CERCANO', existencia=3, dias=30)
        segundo = self.crear_lote('LOT-LEJANO', existencia=10, dias=300)
        with transaction.atomic():
            lote_representativo, costo, asignaciones = surtir_fefo(self.medicamento.id, 5)
        primero.refresh_from_db()
        segundo.refresh_from_db()
        self.assertEqual(lote_representativo, primero)
        self.assertEqual((primero.existencia, segundo.existencia), (0, 8))
        self.assertEqual(sum(a['cantidad'] for a in asignaciones), 5)
        self.assertEqual(costo, primero.costo_unitario * 5)

    def test_folios_automaticos_son_unicos_y_consecutivos(self):
        base = {
            'tipo_entrada': 'ALMACEN', 'almacen': self.almacen,
            'fuente_financiamiento': self.fuente, 'proceso': 'Compra',
            'recibido_por': self.usuario,
        }
        primera = Entrada.objects.create(**base)
        segunda = Entrada.objects.create(**base)
        self.assertNotEqual(primera.folio, segunda.folio)
        self.assertTrue(primera.folio.endswith('-0001'))
        self.assertTrue(segunda.folio.endswith('-0002'))
        self.assertEqual(
            FolioConsecutivo.objects.get(tipo='ENT').ultimo_numero,
            2,
        )

    def test_base_de_datos_impide_existencia_negativa(self):
        lote = self.crear_lote(existencia=1)
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Lote.objects.filter(pk=lote.pk).update(existencia=-1)

    def test_entrada_usa_al_usuario_autenticado_como_responsable(self):
        conceder(self.usuario, 'add_entrada')
        otro = UsuarioPersonalizado.objects.create_user(
            username='otro', password='PruebaSegura123!', rol='FARMACIA'
        )
        self.client.force_login(self.usuario)
        payload = {
            'folio': 'ENT-PRUEBA-0001', 'fecha': date.today().isoformat(),
            'tipo_entrada': 'ALMACEN', 'almacen': self.almacen.id,
            'fuente_financiamiento': self.fuente.id, 'proceso': 'Compra',
            'recibido_por': otro.id,
            'detalles': [{
                'medicamento_id': self.medicamento.id, 'lote': 'ENT-L001',
                'caducidad': (date.today() + timedelta(days=300)).isoformat(),
                'cantidad': 5, 'precio_unitario': '12.50',
                'presentacion_id': self.presentacion.id,
            }],
        }
        response = self.client.post(
            reverse('guardar_entradas'), data=json.dumps(payload),
            content_type='application/json'
        )
        self.assertEqual(response.status_code, 200, response.content)
        entrada = Entrada.objects.get(folio='ENT-PRUEBA-0001')
        self.assertEqual(entrada.recibido_por, self.usuario)
        self.assertEqual(DetalleEntrada.objects.filter(entrada=entrada).count(), 1)

    def test_salida_registra_receta_y_descuenta_stock(self):
        conceder(self.usuario, 'create_salida')
        lote = self.crear_lote(existencia=10)
        self.client.force_login(self.usuario)
        response = self.client.post(reverse('registrar_salida'), {
            'paciente_curp': 'SALIPRUEBA12345678', 'paciente_nombre': 'Paciente Salida',
            'paciente_nacimiento': '1990-01-01', 'receta_origen': 'urgencias',
            'receta_folio': 'REC-PRUEBA-0001', 'item_lote_0': lote.id,
            'item_cantidad_0': '4',
        })
        self.assertEqual(response.status_code, 200, response.content)
        lote.refresh_from_db()
        self.assertEqual(lote.existencia, 6)
        receta = Receta.objects.get(id_folio='REC-PRUEBA-0001')
        self.assertEqual(receta.estado, 'completa')
        self.assertEqual(receta.medicamentos_no_surtidos.count(), 0)

    def test_salida_parcial_guarda_faltante_asociado_a_la_receta(self):
        conceder(self.usuario, 'create_salida')
        lote = self.crear_lote(existencia=10)
        self.client.force_login(self.usuario)

        response = self.client.post(reverse('registrar_salida'), {
            'paciente_curp': 'PARCPRUEBA12345678',
            'paciente_nombre': 'Paciente Parcial',
            'paciente_nacimiento': '1990-01-01',
            'receta_origen': 'urgencias',
            'receta_folio': 'REC-PARCIAL-0001',
            'item_lote_0': lote.id,
            'item_cantidad_0': '4',
            'faltante_desc_0': 'Medicamento sin existencia',
            'faltante_cant_0': '7',
            'faltante_motivo_0': 'Sin stock disponible',
        })

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['estado'], 'parcial')
        receta = Receta.objects.get(id_folio='REC-PARCIAL-0001')
        self.assertEqual(receta.estado, 'parcial')
        self.assertEqual(receta.recetamedicamento_set.count(), 1)
        faltante = receta.medicamentos_no_surtidos.get()
        self.assertEqual(faltante.medicamento_descripcion, 'Medicamento sin existencia')
        self.assertEqual(faltante.cantidad_solicitada, 7)
        self.assertEqual(faltante.motivo, 'Sin stock disponible')
        self.assertEqual(faltante.registrado_por, self.usuario)

    def test_salida_no_surtida_conserva_todos_los_faltantes_en_su_receta(self):
        conceder(self.usuario, 'create_salida')
        self.client.force_login(self.usuario)

        response = self.client.post(reverse('registrar_salida'), {
            'paciente_curp': 'NOSUPRUEBA1234567',
            'paciente_nombre': 'Paciente No Surtido',
            'paciente_nacimiento': '1990-01-01',
            'receta_origen': 'consulta_externa',
            'receta_folio': 'REC-NOSURT-0001',
            'faltante_desc_0': 'Medicamento A',
            'faltante_cant_0': '2',
            'faltante_motivo_0': 'Sin stock disponible',
            'faltante_desc_1': 'Medicamento B',
            'faltante_cant_1': '3',
            'faltante_motivo_1': 'Medicamento no existe en inventario',
        })

        self.assertEqual(response.status_code, 200, response.content)
        receta = Receta.objects.get(id_folio='REC-NOSURT-0001')
        self.assertEqual(receta.estado, 'no_surtida')
        self.assertEqual(receta.recetamedicamento_set.count(), 0)
        self.assertEqual(receta.medicamentos_no_surtidos.count(), 2)

    def test_faltante_invalido_no_descuenta_stock_ni_crea_receta(self):
        conceder(self.usuario, 'create_salida')
        lote = self.crear_lote(existencia=10)
        self.client.force_login(self.usuario)

        response = self.client.post(reverse('registrar_salida'), {
            'paciente_curp': 'INVAPRUEBA12345678',
            'paciente_nombre': 'Paciente Inválido',
            'paciente_nacimiento': '1990-01-01',
            'receta_origen': 'urgencias',
            'receta_folio': 'REC-INVALIDA-001',
            'item_lote_0': lote.id,
            'item_cantidad_0': '4',
            'faltante_desc_0': 'Medicamento inválido',
            'faltante_cant_0': '0',
            'faltante_motivo_0': 'Sin stock disponible',
        })

        self.assertEqual(response.status_code, 400, response.content)
        lote.refresh_from_db()
        self.assertEqual(lote.existencia, 10)
        self.assertFalse(Receta.objects.filter(id_folio='REC-INVALIDA-001').exists())
        self.assertEqual(MedicamentoNoSurtido.objects.count(), 0)

    def test_transferencia_descuenta_y_registra_detalle(self):
        conceder(self.usuario, 'create_transferencia')
        lote = self.crear_lote(existencia=8)
        institucion = Institucion.objects.create(
            codigo='INST-01', nombre='Institución destino', tipo='HOSPITAL'
        )
        self.client.force_login(self.usuario)
        response = self.client.post(reverse('registrar_salida_transferencia'), {
            'institucion_destino': institucion.id, 'item_lote_0': lote.id,
            'item_cantidad_0': '3',
        })
        self.assertEqual(response.status_code, 200, response.content)
        lote.refresh_from_db()
        self.assertEqual(lote.existencia, 5)
        self.assertEqual(DetalleSalidaTransferencia.objects.count(), 1)

    def test_carga_masiva_exige_permiso_de_subida_y_procesa_excel(self):
        def crear_archivo():
            wb = Workbook()
            ws = wb.active
            ws.append(['clave', 'descripcion', 'lote', 'cantidad', 'precio', 'caducidad',
                       'origen', 'contrato', 'fuente_financiamiento'])
            ws.append(['010.000.0099', 'Medicamento carga masiva', 'MAS-001', 5, 10,
                       date.today() + timedelta(days=400), 'ALMACEN', 'CONT-1', 'FEDERAL'])
            archivo = BytesIO()
            wb.save(archivo)
            archivo.seek(0)
            archivo.name = 'carga.xlsx'
            return archivo

        self.client.force_login(self.usuario)
        self.assertEqual(
            self.client.post(reverse('procesar_carga_masiva'), {'archivo': crear_archivo()}).status_code,
            403
        )
        conceder(self.usuario, 'upload_carga_masiva')
        response = self.client.post(
            reverse('procesar_carga_masiva'), {'archivo': crear_archivo()}
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(Medicamento.objects.filter(clave='010.000.0099').exists())


class RegistroRecetasReportesTests(DatosFarmaciaMixin, TestCase):
    def setUp(self):
        self.crear_catalogos()
        self.lote = self.crear_lote(existencia=20)
        self.admin = UsuarioPersonalizado.objects.create_superuser(
            username='admin-reportes-recetas',
            password='PruebaSegura123!',
            rol='ADMIN',
            first_name='Ana',
            last_name='Farmacia',
        )
        self.paciente = Paciente.objects.create(
            nombre_completo='Paciente Reporte',
            curp='REPOPRUEBA12345678',
            fecha_nacimiento=date(1990, 1, 1),
        )
        self.client.force_login(self.admin)

    def crear_receta(self, folio, estado):
        return Receta.objects.create(
            id_folio=folio,
            paciente=self.paciente,
            fecha_emision=date.today(),
            fecha_surtido=date.today(),
            estado=estado,
            origen='urgencias',
            surtido_por=self.admin,
        )

    def test_api_agrupa_surtidos_y_faltantes_por_receta(self):
        receta = self.crear_receta('REC-REPORTE-0001', 'parcial')
        RecetaMedicamento.objects.create(
            receta=receta,
            medicamento=self.medicamento,
            lote=self.lote,
            cantidad_solicitada=4,
            cantidad_surtida=4,
        )
        MedicamentoNoSurtido.objects.create(
            receta=receta,
            medicamento_descripcion='Medicamento faltante',
            cantidad_solicitada=2,
            motivo='Sin stock disponible',
            registrado_por=self.admin,
        )

        response = self.client.get(reverse('api_registro_recetas'), {
            'fecha_inicio': date.today().isoformat(),
            'fecha_fin': date.today().isoformat(),
        })

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual(payload['total_recetas'], 1)
        self.assertEqual(payload['total_renglones'], 2)
        self.assertEqual(payload['data'][0]['folio'], receta.id_folio)
        self.assertEqual(payload['data'][0]['estado'], 'parcial')
        self.assertEqual(
            [item['tipo'] for item in payload['data'][0]['medicamentos']],
            ['surtido', 'no_surtido'],
        )
        self.assertEqual(
            payload['data'][0]['medicamentos'][1]['motivo'],
            'Sin stock disponible',
        )

    def test_api_filtra_estado_y_excluye_colectivos(self):
        self.crear_receta('REC-COMPLETA-001', 'completa')
        parcial = self.crear_receta('REC-PARCIAL-002', 'parcial')
        MedicamentoNoSurtido.objects.create(
            receta=parcial,
            medicamento_descripcion='Medicamento pendiente',
            cantidad_solicitada=1,
            motivo='Sin stock disponible',
        )
        colectivo = self.crear_receta('COL-20260918-0001', 'parcial')
        MedicamentoNoSurtido.objects.create(
            receta=colectivo,
            medicamento_descripcion='Faltante colectivo',
            cantidad_solicitada=1,
            motivo='Sin stock disponible',
        )

        response = self.client.get(reverse('api_registro_recetas'), {
            'fecha_inicio': date.today().isoformat(),
            'fecha_fin': date.today().isoformat(),
            'estado': 'parcial',
        })

        self.assertEqual(response.status_code, 200, response.content)
        folios = [receta['folio'] for receta in response.json()['data']]
        self.assertEqual(folios, ['REC-PARCIAL-002'])

    def test_api_rechaza_filtro_de_estado_invalido(self):
        response = self.client.get(reverse('api_registro_recetas'), {
            'estado': 'estado-inexistente',
        })

        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.json()['success'])

    def test_exportaciones_respetan_el_filtro_de_estado(self):
        completa = self.crear_receta('REC-EXPORT-COMP', 'completa')
        parcial = self.crear_receta('REC-EXPORT-PARC', 'parcial')
        for receta in (completa, parcial):
            RecetaMedicamento.objects.create(
                receta=receta,
                medicamento=self.medicamento,
                lote=self.lote,
                cantidad_solicitada=2,
                cantidad_surtida=2,
            )
        MedicamentoNoSurtido.objects.create(
            receta=parcial,
            medicamento_descripcion='Medicamento pendiente para exportar',
            cantidad_solicitada=3,
            motivo='Sin stock disponible',
        )
        parametros = {
            'fecha_inicio': date.today().isoformat(),
            'fecha_fin': date.today().isoformat(),
            'estado': 'parcial',
        }

        excel = self.client.get(reverse('exportar_registro_recetas_excel'), parametros)
        self.assertEqual(excel.status_code, 200, excel.content[:500])
        self.assertEqual(
            excel['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        workbook = load_workbook(BytesIO(excel.content), read_only=True)
        worksheet = workbook['Registro de Recetas']
        contenido = '\n'.join(
            str(valor)
            for fila in worksheet.iter_rows(values_only=True)
            for valor in fila
            if valor is not None
        )
        self.assertIn('REC-EXPORT-PARC', contenido)
        self.assertIn('Medicamento pendiente para exportar', contenido)
        self.assertNotIn('REC-EXPORT-COMP', contenido)

        pdf = self.client.get(reverse('exportar_registro_recetas_pdf'), parametros)
        self.assertEqual(pdf.status_code, 200, pdf.content[:500])
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF'))
        self.assertIn('Registro_Recetas_parcial_', pdf['Content-Disposition'])


@skipUnlessDBFeature('has_select_for_update')
class ConcurrenciaInventarioTests(DatosFarmaciaMixin, TransactionTestCase):
    reset_sequences = True

    def setUp(self):
        self.crear_catalogos()
        self.lote = self.crear_lote(existencia=10)

    def test_dos_descuentos_simultaneos_no_generan_stock_negativo(self):
        barrera = threading.Barrier(2)
        resultados = []

        def descontar():
            close_old_connections()
            try:
                barrera.wait(timeout=5)
                with transaction.atomic():
                    descontar_lotes({self.lote.id: 6})
                resultados.append('ok')
            except ValidationError:
                resultados.append('sin_stock')
            finally:
                close_old_connections()

        hilos = [threading.Thread(target=descontar) for _ in range(2)]
        for hilo in hilos:
            hilo.start()
        for hilo in hilos:
            hilo.join(timeout=10)
        self.lote.refresh_from_db()
        self.assertEqual(sorted(resultados), ['ok', 'sin_stock'])
        self.assertEqual(self.lote.existencia, 4)
